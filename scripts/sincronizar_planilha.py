#!/usr/bin/env python3
import copy
import datetime as dt
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

WORKBOOK = Path("Radar_Brasil_2027.xlsx")
EVENTS_JSON = Path("dados.json")
NEWS_JSON = Path("noticias.json")
EVENTS_SHEET = "02_Eventos"
NEWS_SHEET = "06_Noticias"

EVENT_FIELDS = [
    "ID", "Titulo", "Status", "Data", "DataBR", "DataCalendario", "DataFim", "Hora", "HoraFim", "FusoHorario",
    "UF", "Cidade", "Categoria", "Organizador", "Publico", "Patrocinador", "Local",
    "Latitude", "Longitude", "Link", "Observacoes", "Mes", "Ano", "Regiao",
]
EVENT_OPTIONAL_FIELDS = {"Titulo", "DataBR", "DataCalendario", "DataFim", "Mes", "Ano", "Regiao", "Hora", "HoraFim", "FusoHorario"}
EVENT_TIME_FIELDS = ("DataCalendario", "DataFim", "Hora", "HoraFim", "FusoHorario")
NEWS_FIELDS = [
    "Data", "Titulo", "Tema", "CidadeUF", "Veiculo", "Link",
    "Sentimento", "Impacto", "Resumo",
]

ALIASES = {
    "publicoestimado": "Publico",
    "linkfonte": "Link",
}


def normalized(value):
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text.lower())


def load_json(path):
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise RuntimeError(f"{path} não contém uma lista JSON.")
    return data


def as_excel_date(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    return value


def find_header(ws, fields, max_rows=50):
    wanted = {normalized(field): field for field in fields}
    wanted.update(ALIASES)
    best = None
    for row_num in range(1, min(ws.max_row, max_rows) + 1):
        found = {}
        for col_num in range(1, ws.max_column + 1):
            key = normalized(ws.cell(row_num, col_num).value)
            canonical = wanted.get(key)
            if canonical:
                found[canonical] = col_num
        if best is None or len(found) > len(best[1]):
            best = (row_num, found)
        required = set(fields)
        if fields == EVENT_FIELDS:
            required -= EVENT_OPTIONAL_FIELDS
        if required.issubset(found):
            return row_num, found
    if best:
        missing = [field for field in fields if field not in best[1]]
        raise RuntimeError(
            f"Cabeçalho de {ws.title} não reconhecido. Campos não localizados: {', '.join(missing)}"
        )
    raise RuntimeError(f"Cabeçalho de {ws.title} não encontrado.")


def copy_row_style(ws, source_row, target_row):
    if source_row < 1 or source_row > ws.max_row or source_row == target_row:
        return
    for col in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.protection:
            dst.protection = copy.copy(src.protection)
    if source_row in ws.row_dimensions:
        src_dim = ws.row_dimensions[source_row]
        dst_dim = ws.row_dimensions[target_row]
        dst_dim.height = src_dim.height
        dst_dim.hidden = src_dim.hidden
        dst_dim.outlineLevel = src_dim.outlineLevel


def update_filters_and_tables(ws, header_row, last_row, data_max_col=None):
    if ws.auto_filter and ws.auto_filter.ref:
        min_col, min_row, max_col, _ = range_boundaries(ws.auto_filter.ref)
        if min_row == header_row:
            max_col = max(max_col, data_max_col or max_col)
            ws.auto_filter.ref = (
                f"{get_column_letter(min_col)}{header_row}:"
                f"{get_column_letter(max_col)}{last_row}"
            )
    for table in ws.tables.values():
        min_col, min_row, max_col, _ = range_boundaries(table.ref)
        if min_row == header_row:
            max_col = max(max_col, data_max_col or max_col)
            table.ref = (
                f"{get_column_letter(min_col)}{header_row}:"
                f"{get_column_letter(max_col)}{last_row}"
            )


def sync_sheet(ws, rows, fields):
    header_row, columns = find_header(ws, fields)

    # A planilha histórica não possuía colunas de horário. Criamos somente as
    # colunas opcionais de tempo quando ausentes, preservando todas as demais.
    if fields == EVENT_FIELDS:
        next_col = max(ws.max_column, max(columns.values(), default=0)) + 1
        style_source_col = max(columns.values(), default=1)
        for field in EVENT_TIME_FIELDS:
            if field in columns:
                continue
            col_num = next_col
            next_col += 1
            columns[field] = col_num
            src = ws.cell(header_row, style_source_col)
            dst = ws.cell(header_row, col_num)
            dst.value = field
            if src.has_style:
                dst._style = copy.copy(src._style)
            if src.alignment:
                dst.alignment = copy.copy(src.alignment)
            if src.protection:
                dst.protection = copy.copy(src.protection)

    first_data_row = header_row + 1
    old_last_row = ws.max_row
    template_row = first_data_row if first_data_row <= old_last_row else header_row

    # Limpa somente as colunas de dados reconhecidas, preservando o restante da aba.
    for row_num in range(first_data_row, old_last_row + 1):
        for col_num in set(columns.values()):
            ws.cell(row_num, col_num).value = None

    for offset, item in enumerate(rows):
        row_num = first_data_row + offset
        if row_num > old_last_row:
            copy_row_style(ws, template_row, row_num)
        for field, col_num in columns.items():
            value = item.get(field)
            if field in ("Data", "DataCalendario", "DataFim"):
                value = as_excel_date(value)
            ws.cell(row_num, col_num).value = value
        if "Data" in columns:
            ws.cell(row_num, columns["Data"]).number_format = "dd/mm/yyyy"
        if "DataCalendario" in columns:
            ws.cell(row_num, columns["DataCalendario"]).number_format = "dd/mm/yyyy"
        if "DataFim" in columns:
            ws.cell(row_num, columns["DataFim"]).number_format = "dd/mm/yyyy"

    last_row = max(header_row + len(rows), header_row + 1)
    update_filters_and_tables(
        ws, header_row, last_row, data_max_col=max(columns.values(), default=ws.max_column)
    )
    return len(rows)


def main():
    for path in (WORKBOOK, EVENTS_JSON, NEWS_JSON):
        if not path.is_file():
            raise RuntimeError(f"Arquivo obrigatório ausente: {path}")

    events = load_json(EVENTS_JSON)
    news = load_json(NEWS_JSON)
    if not events:
        raise RuntimeError("dados.json está vazio; a planilha não será sobrescrita.")

    wb = load_workbook(WORKBOOK)
    if EVENTS_SHEET not in wb.sheetnames:
        raise RuntimeError(f"Aba ausente: {EVENTS_SHEET}")
    if NEWS_SHEET not in wb.sheetnames:
        raise RuntimeError(f"Aba ausente: {NEWS_SHEET}")

    event_count = sync_sheet(wb[EVENTS_SHEET], events, EVENT_FIELDS)
    news_count = sync_sheet(wb[NEWS_SHEET], news, NEWS_FIELDS)
    wb.save(WORKBOOK)
    print(
        f"Planilha sincronizada: {event_count} eventos e {news_count} notícias "
        f"em {WORKBOOK}."
    )


if __name__ == "__main__":
    main()
